
from pathlib import Path
import csv
import io

import pymupdf
import pandas as pd
from docx import Document as WordDocument
from PIL import Image


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".xlsx",
    ".xls",
    ".csv",
    ".png",
    ".jpg",
    ".jpeg",
}


def extract_document_text(file_bytes, filename):
    """
    Extract text from supported document formats.
    Returns extracted text as a string.
    """
    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: {extension}"
        )

    if extension == ".pdf":
        return extract_pdf_text(file_bytes)

    if extension == ".docx":
        return extract_docx_text(file_bytes)

    if extension in {".xlsx", ".xls"}:
        return extract_excel_text(file_bytes, extension)

    if extension == ".csv":
        return extract_csv_text(file_bytes)

    if extension in {".png", ".jpg", ".jpeg"}:
        return extract_image_text(file_bytes)

    raise ValueError("Could not determine the document type.")


def extract_pdf_text(file_bytes):
    pages_text = []

    with pymupdf.open(
        stream=file_bytes,
        filetype="pdf",
    ) as pdf:
        if pdf.needs_pass:
            raise ValueError(
                "Password-protected PDFs are not supported."
            )

        for page_number, page in enumerate(pdf, start=1):
            text = page.get_text("text").strip()

            pages_text.append(
                f"Page {page_number}\n{text}"
            )

    return "\n\n".join(pages_text).strip()


def extract_docx_text(file_bytes):
    document = WordDocument(io.BytesIO(file_bytes))
    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            values = [
                cell.text.strip()
                for cell in row.cells
            ]
            parts.append(" | ".join(values))

    return "\n".join(parts).strip()


def extract_excel_text(file_bytes, extension):
    engine = "openpyxl" if extension == ".xlsx" else "xlrd"

    workbook = pd.ExcelFile(
        io.BytesIO(file_bytes),
        engine=engine,
    )

    parts = []

    for sheet_name in workbook.sheet_names:
        dataframe = pd.read_excel(
            workbook,
            sheet_name=sheet_name,
            header=None,
        )

        parts.append(f"Sheet: {sheet_name}")
        parts.append(
            dataframe.fillna("").to_string(
                index=False,
                header=False,
            )
        )

    return "\n\n".join(parts).strip()


def extract_csv_text(file_bytes):
    decoded_text = None

    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            decoded_text = file_bytes.decode(encoding)
            break
        except UnicodeDecodeError:
            continue

    if decoded_text is None:
        raise ValueError(
            "Could not decode this CSV file."
        )

    rows = csv.reader(io.StringIO(decoded_text))

    return "\n".join(
        " | ".join(row)
        for row in rows
    ).strip()


def extract_image_text(file_bytes):
    """
    Use the existing Tesseract OCR configuration.
    """
    import subprocess

    from .ocr_utils import TESSERACT_PATH

    with Image.open(io.BytesIO(file_bytes)) as image:
        image = image.convert("RGB")

        buffer = io.BytesIO()
        image.save(buffer, format="PNG")
        image_bytes = buffer.getvalue()

    if not TESSERACT_PATH.is_file():
        raise FileNotFoundError(
            f"Tesseract executable not found: {TESSERACT_PATH}"
        )

    result = subprocess.run(
        [
            str(TESSERACT_PATH),
            "stdin",
            "stdout",
            "-l",
            "tam+eng",
        ],
        input=image_bytes,
        capture_output=True,
        timeout=60,
        check=False,
    )

    if result.returncode != 0:
        error_message = result.stderr.decode(
            "utf-8",
            errors="replace",
        )
        raise RuntimeError(
            f"Image OCR failed: {error_message}"
        )

    return result.stdout.decode(
        "utf-8",
        errors="replace",
    ).strip()
